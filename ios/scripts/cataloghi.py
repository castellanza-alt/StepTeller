#!/usr/bin/env python3
"""Generates the NEAT String Catalogs from the translation sheet.

Source: localizzazione/NEAT_testi_tradotti.xlsx, sheet «Testi» (one row per text, stable ID, one column
per language). Output (never edit by hand: regenerated):
  App/Localizable.xcstrings       app texts
  Shared/Widget.xcstrings         widget texts (table «Widget», in the app and in the extension)
  App/InfoPlist.xcstrings         Health permissions of the app
  LiveActivity/InfoPlist.xcstrings  Health permission of the widget

Rules:
  - the key is the sheet ID, the comment is the sheet «Contesto»;
  - {placeholder} → %@ (several placeholders: %1$@, %2$@… in the order of PARAMS, which is the code order);
  - .one/.other rows → one key with plural variants (%lld); in «one» the written 1 becomes the number;
  - sentences made of several rows (notification with warning, widget VoiceOver): one string per language.

Usage: python3 scripts/cataloghi.py   (needs openpyxl: pip install openpyxl)
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SHEET = ROOT / "localizzazione" / "NEAT_testi_tradotti.xlsx"

# sheet column → language code (the first is the development language)
LANGUAGES = [
    ("Italiano", "it"),
    ("Inglese (EN)", "en"),
    ("Spagnolo (ES)", "es"),
    ("Francese (FR)", "fr"),
    ("Tedesco (DE)", "de"),
    ("Portoghese Brasile (PT-BR)", "pt-BR"),
    ("Giapponese (JA)", "ja"),
    ("Coreano (KO)", "ko"),
    ("Cinese semplificato (ZH-Hans)", "zh-Hans"),
    ("Cinese tradizionale (ZH-Hant)", "zh-Hant"),
]
# plural categories used by iOS (CLDR) per language
PLURAL_CATEGORIES = {"ja": ["other"], "ko": ["other"], "zh-Hans": ["other"], "zh-Hant": ["other"]}

# Parameter order as the code passes them (`L10n.format`). Keys without an entry: no parameters.
PARAMS = {
    "intro.page.vo": ["n", "tot"],
    "home.steps.of": ["obiettivo"],
    "home.steps.vo": ["passi", "obiettivo"],
    "home.done.extra": ["passi"],
    "home.minutes.km": ["km"],
    "home.minutes.vo": ["min", "velocita", "km"],
    "home.speed.value.vo": ["velocita"],
    "top.jolly.vo": ["n"],
    "st.longest.vo": ["n"],
    "st.nextJolly.vo": ["n", "fatti", "tot"],
    "st.jolly.vo": ["n"],
    "st.big.vo": ["valore", "etichetta"],
    "st.second.vo.avgSteps": ["valore"],
    "st.second.vo.avgKm": ["valore"],
    "st.second.vo.totSteps": ["valore"],
    "st.second.vo.totKm": ["valore"],
    "st.chart.goal.vo": ["valore"],
    "st.chart.bar.vo": ["giorno", "valore"],
    "st.best.vo": ["valore", "unita", "quando"],
    "set.streak.startValue": ["data", "n"],
    "set.version": ["versione"],
    "party.of": ["obiettivo"],
    "party.vo": ["passi", "obiettivo"],
    "w.of": ["obiettivo"],
    "w.minKm": ["km"],
    "w.lock.left": ["min", "km"],
    "w.inline": ["passi"],
    "w.vo.steps": ["passi", "obiettivo"],
    "n.body": ["min", "velocita", "obiettivo", "mancano"],
    # composed
    "n.body.jolly": ["min", "velocita", "obiettivo", "mancano", "n"],
    "n.body.lose": ["min", "velocita", "obiettivo", "mancano", "n"],
    "w.vo.summary.done": ["passi", "obiettivo", "n", "jolly"],
    "w.vo.summary.left": ["passi", "obiettivo", "min", "km", "n", "jolly"],
}

# Whole sentences made of several sheet rows. The separator depends on the language only where the row
# does not carry its own leading punctuation (notification: sentence after sentence).
def sentence_gap(lang):
    return "" if lang in ("ja", "zh-Hans", "zh-Hant") else " "

COMPOSED = {
    "n.body.jolly": (["n.body", "n.jolly"], sentence_gap, "Corpo della notifica con l'avviso sul Jolly (n.body + n.jolly)."),
    "n.body.lose": (["n.body", "n.lose"], sentence_gap, "Corpo della notifica con l'avviso sulla serie (n.body + n.lose)."),
    "w.vo.summary.done": (["w.vo.steps", "w.vo.done", "w.vo.streak"], lambda _: "",
                          "VoiceOver del widget a obiettivo raggiunto (w.vo.steps + w.vo.done + w.vo.streak)."),
    "w.vo.summary.left": (["w.vo.steps", "w.vo.left", "w.vo.streak"], lambda _: "",
                          "VoiceOver del widget con minuti e km che mancano (w.vo.steps + w.vo.left + w.vo.streak)."),
}
# rows used only inside composed sentences
ONLY_COMPOSED = {"n.jolly", "n.lose", "w.vo.done", "w.vo.left", "w.vo.streak"}

INFO_PLIST = {
    "p.health.read": ("app", "NSHealthShareUsageDescription"),
    "p.health.write": ("app", "NSHealthUpdateUsageDescription"),
    "p.health.widget": ("widget", "NSHealthShareUsageDescription"),
}


def load_rows():
    ws = openpyxl.load_workbook(SHEET)["Testi"]
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    out = {}
    for r in rows[1:]:
        if not r[0]:
            continue
        row = dict(zip(header, r))
        for col, lang in LANGUAGES:
            if not (row.get(col) or "").strip():
                sys.exit(f"{row['ID']}: manca la traduzione {lang}")
        out[row["ID"]] = row
    return out


def placeholders(text):
    return re.findall(r"\{(\w+)\}", text)


def to_format(key, text, params, lang):
    """{name} → %@ / %N$@ following the order of PARAMS."""
    found = placeholders(text)
    if sorted(found) != sorted(params):
        sys.exit(f"{key} [{lang}]: segnaposto {found}, attesi {params}")
    if params:
        text = text.replace("%", "%%")
    for name in params:
        spec = "%@" if len(params) == 1 else f"%{params.index(name) + 1}$@"
        text = text.replace("{" + name + "}", spec)
    return text


def unit(value):
    return {"stringUnit": {"state": "translated", "value": value}}


def plural_value(key, one, other, lang):
    """Plural variants on the first argument (%lld). In «one» the written 1 becomes the number."""
    def arg(text, form):
        if "{n}" in text:
            if placeholders(text) != ["n"]:
                sys.exit(f"{key}.{form} [{lang}]: segnaposto inattesi {placeholders(text)}")
            return text.replace("%", "%%").replace("{n}", "%arg")
        if placeholders(text):
            sys.exit(f"{key}.{form} [{lang}]: segnaposto inattesi {placeholders(text)}")
        ones = re.findall(r"(?<!\d)1(?!\d)", text)
        if form == "one" and len(ones) == 1:
            return re.sub(r"(?<!\d)1(?!\d)", "%arg", text.replace("%", "%%"))
        if len(ones) > 1:
            sys.exit(f"{key}.{form} [{lang}]: più di un 1 nel testo")
        return text.replace("%", "%%")

    forms = {"one": one, "other": other}
    cats = PLURAL_CATEGORIES.get(lang, ["one", "other"])
    variations = {c: unit(arg(forms[c], c)) for c in cats}
    return {
        "stringUnit": {"state": "translated", "value": "%#@n@"},
        "substitutions": {"n": {"argNum": 1, "formatSpecifier": "lld", "variations": {"plural": variations}}},
    }


def comment(row):
    limit = row.get("Limite (caratteri)") or "—"
    ctx = (row.get("Contesto") or "").strip()
    return ctx if limit in ("—", None) else f"{ctx} Limite: {limit}."


def build():
    rows = load_rows()
    app, widget, info_app, info_widget = {}, {}, {}, {}
    used = set()                    # rows already placed (plurals and permissions)

    def target(key):
        return widget if key.startswith("w.") else app

    # plurals
    for key in sorted({k.rsplit(".", 1)[0] for k in rows if k.endswith((".one", ".other"))}):
        one, other = rows[f"{key}.one"], rows[f"{key}.other"]
        used |= {f"{key}.one", f"{key}.other"}
        target(key)[key] = {
            "comment": f"{comment(one)} Plurale ({key}.one / {key}.other).",
            "extractionState": "manual",
            "localizations": {lang: plural_value(key, one[col], other[col], lang) for col, lang in LANGUAGES},
        }
    # composed sentences
    for key, (parts, gap, note) in COMPOSED.items():
        locs = {}
        for col, lang in LANGUAGES:
            sep = gap(lang)
            text = sep.join(rows[p][col].strip() if sep else rows[p][col] for p in parts)
            locs[lang] = unit(to_format(key, text, PARAMS[key], lang))
        target(key)[key] = {"comment": note, "extractionState": "manual", "localizations": locs}
    # permissions
    for key, (where, plist_key) in INFO_PLIST.items():
        used.add(key)
        row = rows[key]
        table = info_app if where == "app" else info_widget
        table[plist_key] = {
            "comment": f"{key}: {comment(row)}",
            "extractionState": "manual",
            "localizations": {lang: unit(row[col].strip()) for col, lang in LANGUAGES},
        }
    # all the other rows
    for key, row in rows.items():
        if key in used or key in ONLY_COMPOSED:
            continue
        params = PARAMS.get(key, [])
        target(key)[key] = {
            "comment": comment(row),
            "extractionState": "manual",
            "localizations": {lang: unit(to_format(key, row[col], params, lang)) for col, lang in LANGUAGES},
        }
    missing = set(PARAMS) - set(app) - set(widget)
    if missing:
        sys.exit(f"chiavi in PARAMS senza riga: {sorted(missing)}")
    return {"App/Localizable.xcstrings": app, "Shared/Widget.xcstrings": widget,
            "App/InfoPlist.xcstrings": info_app, "LiveActivity/InfoPlist.xcstrings": info_widget}


def write(path, strings):
    catalog = {"sourceLanguage": "it", "strings": strings, "version": "1.0"}
    # same format Xcode uses to save catalogs, so it does not rewrite them on opening
    text = json.dumps(catalog, ensure_ascii=False, indent=2, sort_keys=True, separators=(",", " : "))
    (ROOT / path).write_text(text + "\n", encoding="utf-8")
    print(f"{path}: {len(strings)} chiavi")


if __name__ == "__main__":
    for path, strings in build().items():
        write(path, strings)
