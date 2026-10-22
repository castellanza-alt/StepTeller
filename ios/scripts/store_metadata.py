#!/usr/bin/env python3
"""Exports the App Store texts from the sheets to store/metadata.json (read by `asc.py store-*`).

Sources (only read; the JSON is regenerated, never edited by hand: CI reads it so that it needs no spreadsheet library):
  localizzazione/NEAT_store_tradotto.xlsx, sheet «Store»: rows store.name, store.subtitle, store.promo,
    store.description, store.keywords (Italian + 9 languages) → "languages";
  localizzazione/NEAT_novita.xlsx, sheet «Novità»: one row per version (novita.1.1) plus the fixed novita.correzioni
    row for fix-only versions → "whatsNew" (≤ 4000 characters per language, no empty cell). See RILASCIO.md.

Keywords: Apple documents the field as «up to 100 bytes» (UTF-8). The Japanese, Korean and Chinese rows of the sheet
are 102–146 bytes, so KEYWORD_OVERRIDES holds shorter versions proposed to Will (whole words removed, none cut).

Usage: python3 scripts/store_metadata.py   (needs openpyxl: pip install openpyxl)
"""
import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SHEET = ROOT / "localizzazione" / "NEAT_store_tradotto.xlsx"
NEWS_SHEET = ROOT / "localizzazione" / "NEAT_novita.xlsx"
OUT = ROOT / "store" / "metadata.json"

# sheet column → language code (same codes as the screenshot folders)
LANGUAGES = [
    ("Italiano", "it"),
    ("Inglese (EN)", "en"),
    ("Spagnolo (ES)", "es"),
    ("Francese (FR)", "fr"),
    ("Tedesco (DE)", "de"),
    ("Portoghese Brasile (PT-BR)", "pt-BR"),
    ("Giapponese (JA)", "ja"),
    ("Coreano (KO)", "ko"),
    ("Cinese semplificato (ZH-Hans)", "zh-Hans"),
    ("Cinese tradizionale (ZH-Hant)", "zh-Hant"),
]

# sheet ID → field name in the JSON
FIELDS = {
    "store.name": "name",
    "store.subtitle": "subtitle",
    "store.promo": "promotionalText",
    "store.description": "description",
    "store.keywords": "keywords",
}

# Shortened keywords (≤ 100 bytes), whole words dropped from the sheet row; pending Will's approval.
KEYWORD_OVERRIDES = {
    # dropped: カウンター, 徒歩, 時間, リマインダー
    "ja": "歩数計,ウォーキング,散歩,歩数,距離,記録,連続達成,ウィジェット,10000歩",
    # dropped: 측정
    "ko": "만보기,걷기,산책,걸음수,거리,시간,기록,위젯,알림,만보,연속달성,일일",
    # dropped: 计步 (already in 计步器), 步行时间, 步行距离, 步数记录
    "zh-Hans": "计步器,步数,走路,散步,步行,万步,每日步数,连续达标,小组件,提醒,10000步",
    # dropped: 計步 (already in 計步器), 步行時間, 步行距離, 步數紀錄
    "zh-Hant": "計步器,步數,走路,健走,散步,萬步,每日步數,連續達標,小工具,提醒,10000步",
}

LIMITS = {"name": 30, "subtitle": 30, "promotionalText": 170, "description": 4000}
KEYWORD_BYTES = 100
WHATS_NEW_LIMIT = 4000
WHATS_NEW_PREFIX = "novita."
VERSION_RE = re.compile(r"^\d+\.\d+(\.\d+)?$")


def read_whats_new():
    """Sheet «Novità» → {version or "correzioni": {language: text}}; stops on empty cells, bad IDs or long texts."""
    sheet = openpyxl.load_workbook(NEWS_SHEET, read_only=True, data_only=True)["Novità"]
    rows = list(sheet.iter_rows(values_only=True))
    header = [str(c).strip() if c else "" for c in rows[0]]
    columns = {code: header.index(title) for title, code in LANGUAGES}
    news, problems = {}, []
    for row in rows[1:]:
        identifier = str(row[0] or "").strip()
        if not identifier:
            continue
        if not identifier.startswith(WHATS_NEW_PREFIX):
            problems.append(f"Novità: ID «{identifier}» non inizia con {WHATS_NEW_PREFIX}")
            continue
        key = identifier[len(WHATS_NEW_PREFIX):]
        if key != "correzioni" and not VERSION_RE.match(key):
            problems.append(f"Novità: ID «{identifier}» non è novita.<versione> (es. novita.1.1) né novita.correzioni")
            continue
        if key in news:
            problems.append(f"Novità: riga {identifier} ripetuta")
            continue
        texts = {}
        for code, index in columns.items():
            value = str(row[index] or "").strip().replace("\r\n", "\n")
            if not value:
                problems.append(f"Novità {identifier} / {code}: cella vuota")
            elif len(value) > WHATS_NEW_LIMIT:
                problems.append(f"Novità {identifier} / {code}: {len(value)} caratteri > {WHATS_NEW_LIMIT}")
            texts[code] = value
        news[key] = texts
    if "correzioni" not in news:
        problems.append("Novità: manca la riga fissa novita.correzioni")
    if problems:
        raise SystemExit("\n".join(problems))
    return news


def main():
    sheet = openpyxl.load_workbook(SHEET, read_only=True, data_only=True)["Store"]
    rows = list(sheet.iter_rows(values_only=True))
    header = [str(c).strip() if c else "" for c in rows[0]]
    columns = {code: header.index(title) for title, code in LANGUAGES}
    texts = {code: {} for _, code in LANGUAGES}
    for row in rows[1:]:
        field = FIELDS.get(str(row[0] or "").strip())
        if not field:
            continue
        for code, index in columns.items():
            value = str(row[index] or "").strip()
            if not value:
                raise SystemExit(f"Cella vuota: {row[0]} / {code}")
            texts[code][field] = value
    for code, keywords in KEYWORD_OVERRIDES.items():
        texts[code]["keywords"] = keywords
    problems = []
    for code, fields in texts.items():
        missing = set(FIELDS.values()) - set(fields)
        if missing:
            problems.append(f"{code}: mancano {sorted(missing)}")
        for field, limit in LIMITS.items():
            if len(fields.get(field, "")) > limit:
                problems.append(f"{code} {field}: {len(fields[field])} caratteri > {limit}")
        size = len(fields.get("keywords", "").encode())
        if size > KEYWORD_BYTES:
            problems.append(f"{code} keywords: {size} byte > {KEYWORD_BYTES}")
    if problems:
        raise SystemExit("\n".join(problems))
    news = read_whats_new()
    OUT.parent.mkdir(exist_ok=True)
    data = {"_source": "Generato da scripts/store_metadata.py dai fogli «Store» e «Novità»: non modificare a mano.",
            "languages": texts, "whatsNew": news}
    OUT.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    versions = ", ".join(k for k in news if k != "correzioni") or "nessuna versione"
    print(f"{OUT.relative_to(ROOT)}: {len(texts)} lingue; Novità: {versions} + correzioni")
    for code, fields in texts.items():
        print(f"  {code:8} nome {len(fields['name']):2}  sottotitolo {len(fields['subtitle']):2}  "
              f"promo {len(fields['promotionalText']):3}  descrizione {len(fields['description']):4}  "
              f"parole chiave {len(fields['keywords']):3} car. / {len(fields['keywords'].encode()):3} byte")


if __name__ == "__main__":
    main()
